"""Build the compilation documents (DOCX + PDF) of Mangalili, Fhetz.

Run from the repository root:
    python3 tools/build_documents.py

Reads the source code inside COMPILATION/, renders one figure per form design
and writes:
    docs/Mangalili_Fhetz_Activity_Compilation.docx
    docs/Mangalili_Fhetz_Activity_Compilation.pdf
    docs/figures/*.png
"""
import os
import re
import sys
import textwrap

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from activities import (ACTIVITIES, AUTHOR, COURSE, SUBJECT, INSTRUCTOR_LINE,
                        SCHOOL_YEAR, DATE_LINE)
from design_figure import render as render_figure

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
COMPILATION = os.path.join(ROOT, "COMPILATION")
DOCS = os.path.join(ROOT, "docs")
FIGURES = os.path.join(DOCS, "figures")
DOC_BASENAME = "Mangalili_Fhetz_Activity_Compilation"


# ---------------------------------------------------------------------------
# source code helpers
# ---------------------------------------------------------------------------
def read_source(relative_path):
    path = os.path.join(COMPILATION, relative_path)
    with open(path, "r", encoding="utf-8-sig", errors="replace") as handle:
        text = handle.read()
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", "    ")
    lines = []
    long_line_noted = False
    for line in text.split("\n"):
        if len(line) > 320:
            if not long_line_noted:
                lines.append(line[:120].rstrip() + " ...")
                lines.append("' [very long data line omitted in this document]")
                long_line_noted = True
            continue
        long_line_noted = False
        lines.append(line.rstrip())
    while lines and lines[-1] == "":
        lines.pop()
    return lines


def project_files():
    """Return every source file of the compilation, grouped per project folder."""
    groups = []
    for entry in ACTIVITIES:
        base = os.path.join(COMPILATION, entry["folder"])
        found = []
        for root, dirs, files in os.walk(base):
            dirs[:] = sorted(d for d in dirs if d not in ("bin", "obj", "My Project"))
            for name in sorted(files):
                if name.lower().endswith((".vb", ".vbproj", ".sln")):
                    found.append(os.path.relpath(os.path.join(root, name), COMPILATION))
        groups.append((entry["project"], found))
    return groups


# ---------------------------------------------------------------------------
# figures
# ---------------------------------------------------------------------------
def build_figures():
    os.makedirs(FIGURES, exist_ok=True)
    figures = {}
    index = 0
    for entry in ACTIVITIES:
        rendered = []
        for designer_path, caption in entry["figures"]:
            index += 1
            out_name = "fig_act%02d_%d.png" % (entry["number"], index)
            out_path = os.path.join(FIGURES, out_name)
            full_caption = "Figure %d.%d - %s" % (entry["number"], len(rendered) + 1, caption)
            render_figure(os.path.join(COMPILATION, designer_path), out_path, full_caption)
            rendered.append((out_path, full_caption))
        figures[entry["number"]] = rendered
    return figures


# ---------------------------------------------------------------------------
# formatted text (shared by the DOCX and the PDF writer)
# ---------------------------------------------------------------------------
def code_blocks(entry):
    blocks = []
    for relative_path, label in entry["files"]:
        blocks.append((label, read_source(relative_path)))
    return blocks


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def build_docx(figures, out_path):
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Cm, RGBColor

    doc = Document()

    # page setup -------------------------------------------------------------
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, colour in (("Heading 1", 20, (0x0F, 0x17, 0x2A)),
                               ("Heading 2", 12.5, (0x1E, 0x3A, 0x8A)),
                               ("Heading 3", 11, (0x33, 0x41, 0x55))):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(*colour)
        style.font.bold = True

    def shade(paragraph, fill):
        properties = paragraph._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        properties.append(shd)

    def add_field(paragraph, instruction):
        run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = instruction
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(begin)
        run._r.append(instr)
        run._r.append(end)

    # footer with page number -------------------------------------------------
    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_paragraph.paragraph_format.space_before = Pt(2)
    run = footer_paragraph.add_run("Mangalili, Fhetz  |  " + SUBJECT + "  |  page ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    add_field(footer_paragraph, "PAGE")
    run = footer_paragraph.add_run(" of ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    add_field(footer_paragraph, "NUMPAGES")
    for extra in footer_paragraph.runs:
        extra.font.size = Pt(8)
        extra.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # header -----------------------------------------------------------------
    header_paragraph = section.header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header_paragraph.add_run("Activity Compilation  -  " + SUBJECT)
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    def paragraph(text, size=10.5, bold=False, italic=False, colour=None,
                  align=None, space_after=4, space_before=0, style=None):
        par = doc.add_paragraph(style=style)
        if align == "center":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        par.paragraph_format.space_after = Pt(space_after)
        par.paragraph_format.space_before = Pt(space_before)
        run = par.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.font.name = "Calibri"
        if colour:
            run.font.color.rgb = RGBColor(*colour)
        return par

    def bullet(text):
        par = doc.add_paragraph(text, style="List Bullet")
        par.paragraph_format.space_after = Pt(2)
        for run in par.runs:
            run.font.size = Pt(10.5)
        return par

    def code_block(raw_lines, label=None):
        lines = []
        for line in raw_lines:
            if len(line) > 100:
                lines.extend(textwrap.wrap(line, 100, subsequent_indent="    ",
                                           break_long_words=True,
                                           drop_whitespace=False) or [""])
            else:
                lines.append(line)
        if label:
            label_par = doc.add_paragraph()
            label_par.paragraph_format.space_before = Pt(6)
            label_par.paragraph_format.space_after = Pt(2)
            label_run = label_par.add_run(label)
            label_run.bold = True
            label_run.font.size = Pt(9.5)
            label_run.font.name = "Consolas"
            label_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        for line in lines:
            par = doc.add_paragraph()
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.space_before = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.left_indent = Cm(0.25)
            shade(par, "F2F5F9")
            run = par.add_run(line if line else " ")
            run.font.name = "Consolas"
            run.font.size = Pt(7.5)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ---------------------------- title page --------------------------------
    for _ in range(2):
        doc.add_paragraph()
    paragraph(SUBJECT.upper(), 12, True, colour=(0x37, 0x4F, 0xA0), align="center", space_after=2)
    paragraph("ACTIVITY COMPILATION", 30, True, colour=(0x0F, 0x17, 0x2A), align="center", space_after=6)
    paragraph("Activities 1 to 13  -  Visual Basic .NET Windows Forms and Console Programs",
              11.5, colour=(0x47, 0x55, 0x69), align="center", space_after=26)

    details = [
        ("Submitted by", AUTHOR),
        ("Course, Year and Section", COURSE),
        ("Subject", SUBJECT),
        ("Submitted to", INSTRUCTOR_LINE),
        ("School Year", SCHOOL_YEAR),
        ("Date Submitted", DATE_LINE),
    ]
    for label, value in details:
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(2)
        first = par.add_run(label + " :  ")
        first.bold = True
        first.font.size = Pt(11)
        second = par.add_run(value)
        second.font.size = Pt(11)

    doc.add_paragraph()
    paragraph("This compilation contains the thirteen activities of the semester. "
              "The form designs of Activities 4, 5, 6, 8, 10 and 12 were newly drawn and their "
              "program code was rewritten, while the names of the author, the course and the "
              "program titles were updated in every project.",
              10, italic=True, colour=(0x47, 0x55, 0x69), align="center")

    doc.add_page_break()

    # ---------------------------- contents ----------------------------------
    paragraph("Contents", 18, True, colour=(0x0F, 0x17, 0x2A), space_after=10)
    for entry in ACTIVITIES:
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(3)
        number = par.add_run("Activity %d" % entry["number"])
        number.bold = True
        number.font.size = Pt(10.5)
        title = par.add_run("   " + entry["title"])
        title.font.size = Pt(10.5)
        tail = par.add_run("    (" + entry["project"] + ")")
        tail.font.size = Pt(8.5)
        tail.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    doc.add_paragraph()
    paragraph("Appendix A - Project file list", 11, True, colour=(0x1E, 0x3A, 0x8A), space_before=8)
    paragraph("Appendix B - What changed in this compilation", 11, True, colour=(0x1E, 0x3A, 0x8A), space_after=2)

    # ---------------------------- activities --------------------------------
    for entry in ACTIVITIES:
        doc.add_page_break()
        heading = doc.add_heading(level=1)
        run = heading.add_run("Activity %d" % entry["number"])
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(0x37, 0x4F, 0xA0)
        subtitle = doc.add_paragraph()
        subtitle.paragraph_format.space_after = Pt(2)
        title_run = subtitle.add_run(entry["title"])
        title_run.bold = True
        title_run.font.size = Pt(14)
        title_run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        paragraph("Project : " + entry["project"] + "        Author : " + AUTHOR +
                  "        " + COURSE, 9, colour=(0x64, 0x74, 0x8B), space_after=10)

        paragraph("Description", 12, True, colour=(0x1E, 0x3A, 0x8A), space_before=4, space_after=4)
        for text in entry["description"]:
            paragraph(text, space_after=6)

        if entry.get("objectives"):
            paragraph("Objectives", 12, True, colour=(0x1E, 0x3A, 0x8A),
                      space_before=6, space_after=4)
            for item in entry["objectives"]:
                bullet(item)

        for figure_path, caption in figures.get(entry["number"], []):
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            picture_par = doc.add_paragraph()
            picture_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture_par.paragraph_format.space_after = Pt(2)
            picture_par.add_run().add_picture(figure_path, width=Cm(16.0))
            paragraph(caption, 8.5, italic=True, colour=(0x64, 0x74, 0x8B),
                      align="center", space_after=10)

        paragraph("Source Code", 14, True, colour=(0x1E, 0x3A, 0x8A),
                  space_before=8, space_after=4)
        for label, lines in code_blocks(entry):
            code_block(lines, label)

        if entry.get("output"):
            paragraph("Sample Output", 12, True, colour=(0x1E, 0x3A, 0x8A),
                      space_before=8, space_after=4)
            code_block(entry["output"])

        if entry.get("notes"):
            paragraph("Notes", 12, True, colour=(0x1E, 0x3A, 0x8A),
                      space_before=8, space_after=4)
            for note in entry["notes"]:
                paragraph(note, space_after=4)

    # ---------------------------- appendix A --------------------------------
    doc.add_page_break()
    paragraph("Appendix A - Project file list", 16, True, colour=(0x0F, 0x17, 0x2A), space_after=8)
    paragraph("Every folder below is inside the COMPILATION folder of the submitted archive. "
              "The bin, obj and .suo files of the original projects are not included.",
              9.5, italic=True, colour=(0x47, 0x55, 0x69), space_after=8)
    for project, files in project_files():
        paragraph(project, 10.5, True, colour=(0x1E, 0x3A, 0x8A), space_before=6, space_after=2)
        for item in files:
            paragraph(item.replace("/", "  /  "), 8.5, colour=(0x33, 0x41, 0x55), space_after=0)

    # ---------------------------- appendix B --------------------------------
    doc.add_page_break()
    paragraph("Appendix B - What changed in this compilation", 16, True,
              colour=(0x0F, 0x17, 0x2A), space_after=8)
    changes = [
        ("Author and course", "The name Alvarez, Rex was replaced by Mangalili, Fhetz and the "
                              "course line now reads BSIT 3B in every project, every form title, "
                              "the About form, the application configuration file and the console "
                              "programs."),
        ("Activity 4", "New two panel design with a dark banner, a shared ReadInputs() function, "
                       "four decimal result formatting and a division by zero guard."),
        ("Activity 5", "New side rail layout with three result cards, the Enter key runs the "
                       "evaluation and the first card changes colour with the result."),
        ("Activity 6", "New left card and right result layout with a large average badge, one "
                       "ReadGrades() validation function and a 0 to 100 range check."),
        ("Activity 8", "The four overlapping GroupBoxes were replaced by a compose card and three "
                       "option cards; the long If chain was replaced by helper functions."),
        ("Activity 10", "Both forms were rebuilt: Panel Playground with an event log and an API "
                        "inspector, and Split Container Lab with a category list, a record list, "
                        "a details pane and a reset layout feature."),
        ("Activity 12", "The explorer was rewritten as Mangalili Explorer with history stacks, an "
                        "address bar, a live filter, numeric column sorting, safe image preview and "
                        "layout buttons for the nested splitters."),
    ]
    for label, text in changes:
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(4)
        first = par.add_run(label + "  -  ")
        first.bold = True
        first.font.size = Pt(10.5)
        second = par.add_run(text)
        second.font.size = Pt(10.5)

    doc.add_paragraph()
    paragraph("Prepared by :  " + AUTHOR + "        Course :  " + COURSE +
              "        Date :  " + DATE_LINE, 10, italic=True,
              colour=(0x47, 0x55, 0x69))

    doc.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def build_pdf(figures, out_path):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (BaseDocTemplate, Frame, Image, KeepTogether,
                                    PageBreak, PageTemplate, Paragraph, Preformatted,
                                    Spacer)
    from reportlab.platypus.tableofcontents import TableOfContents

    ink = colors.HexColor("#0F172A")
    navy = colors.HexColor("#1E3A8A")
    blue = colors.HexColor("#374FA0")
    grey = colors.HexColor("#64748B")
    light = colors.HexColor("#94A3B8")
    code_bg = colors.HexColor("#F2F5F9")

    stylesheet = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=stylesheet["BodyText"], fontName="Helvetica",
                          fontSize=9.6, leading=13.4, textColor=ink, alignment=TA_JUSTIFY,
                          spaceAfter=5)
    small = ParagraphStyle("small", parent=body, fontSize=8.2, leading=11, textColor=grey)
    capt = ParagraphStyle("capt", parent=body, fontSize=8, leading=10.5, textColor=grey,
                          alignment=TA_CENTER, spaceBefore=3, spaceAfter=10)
    h_title = ParagraphStyle("htitle", parent=body, fontSize=26, leading=30, textColor=ink,
                             fontName="Helvetica-Bold", alignment=TA_CENTER, spaceAfter=10)
    h_sub = ParagraphStyle("hsub", parent=body, fontSize=11, leading=15, textColor=colors.HexColor("#475569"),
                           alignment=TA_CENTER, spaceAfter=6)
    h_act = ParagraphStyle("hact", parent=body, fontSize=22, leading=26, textColor=blue,
                           fontName="Helvetica-Bold", spaceAfter=2, alignment=0)
    h_act_title = ParagraphStyle("hactt", parent=body, fontSize=13, leading=16, textColor=ink,
                                 fontName="Helvetica-Bold", spaceAfter=2)
    h_sec = ParagraphStyle("hsec", parent=body, fontSize=11.5, leading=14, textColor=navy,
                           fontName="Helvetica-Bold", spaceBefore=8, spaceAfter=4)
    h_small = ParagraphStyle("hsmall", parent=h_sec, fontSize=10)
    bullet_style = ParagraphStyle("bl", parent=body, leftIndent=12, bulletIndent=2, spaceAfter=2)
    code_style = ParagraphStyle("code", parent=stylesheet["Code"], fontName="Courier",
                                fontSize=7.0, leading=8.3, textColor=colors.HexColor("#1F2937"),
                                backColor=code_bg, borderPadding=(3, 4, 3, 4), leftIndent=6,
                                spaceBefore=2, spaceAfter=6, alignment=0)
    code_label = ParagraphStyle("codelabel", parent=body, fontName="Courier-Bold", fontSize=8.4,
                                leading=10.5, textColor=navy, spaceBefore=6, spaceAfter=1)
    toc1 = ParagraphStyle("toc1", parent=body, fontName="Helvetica", fontSize=10, leading=16,
                          alignment=0)
    toc2 = ParagraphStyle("toc2", parent=toc1, fontSize=9, leading=13, leftIndent=14, textColor=grey)
    toc_head = ParagraphStyle("tochead", parent=h_act, fontSize=22, spaceAfter=10)

    class CompilationDoc(BaseDocTemplate):
        def __init__(self, filename, **kwargs):
            BaseDocTemplate.__init__(self, filename, pagesize=A4,
                                     leftMargin=2.0 * cm, rightMargin=1.8 * cm,
                                     topMargin=1.9 * cm, bottomMargin=1.7 * cm,
                                     title="Activity Compilation - " + AUTHOR,
                                     author=AUTHOR, **kwargs)
            frame = Frame(self.leftMargin, self.bottomMargin, self.width, self.height,
                          id="normal")
            self.addPageTemplates([PageTemplate(id="page", frames=[frame],
                                                onPage=self.decorate)])
            self.activity_counter = 0

        def decorate(self, canvas, doc):
            canvas.saveState()
            canvas.setFont("Helvetica", 7.6)
            canvas.setFillColor(light)
            canvas.drawRightString(A4[0] - 1.8 * cm, A4[1] - 1.2 * cm,
                                   "Activity Compilation  -  %s  -  %s" % (SUBJECT, AUTHOR))
            canvas.setStrokeColor(colors.HexColor("#E2E8F0"))
            canvas.setLineWidth(0.6)
            canvas.line(2.0 * cm, A4[1] - 1.35 * cm, A4[0] - 1.8 * cm, A4[1] - 1.35 * cm)
            canvas.setFont("Helvetica", 7.6)
            canvas.setFillColor(grey)
            canvas.drawCentredString(A4[0] / 2.0, 1.05 * cm,
                                     "Mangalili, Fhetz  |  %s  |  page %d" % (SUBJECT, doc.page))
            canvas.restoreState()

        def afterFlowable(self, flowable):
            if isinstance(flowable, Paragraph) and flowable.style.name == "hactt":
                text = flowable.getPlainText()
                self.notify("TOCEntry", (0, text, self.page))

    doc = CompilationDoc(out_path)
    story = []

    # ---------------------------- title page --------------------------------
    story.append(Spacer(1, 3.2 * cm))
    story.append(Paragraph(SUBJECT.upper(), ParagraphStyle("x", parent=h_sub, textColor=blue,
                                                           fontName="Helvetica-Bold",
                                                           fontSize=11, spaceAfter=4)))
    story.append(Paragraph("ACTIVITY COMPILATION", h_title))
    story.append(Paragraph("Activities 1 to 13  -  Visual Basic .NET Windows Forms and "
                           "Console Programs", h_sub))
    story.append(Spacer(1, 1.6 * cm))
    for label, value in (("Submitted by", AUTHOR),
                         ("Course, Year and Section", COURSE),
                         ("Subject", SUBJECT),
                         ("Submitted to", INSTRUCTOR_LINE),
                         ("School Year", SCHOOL_YEAR),
                         ("Date Submitted", DATE_LINE)):
        story.append(Paragraph("<b>%s :</b>  %s" % (label, value), h_sub))
    story.append(Spacer(1, 2.2 * cm))
    story.append(Paragraph("This compilation contains the thirteen activities of the semester. "
                           "The form designs of Activities 4, 5, 6, 8, 10 and 12 were newly drawn "
                           "and their program code was rewritten, while the names of the author, "
                           "the course and the program titles were updated in every project.",
                           ParagraphStyle("y", parent=small, alignment=TA_CENTER,
                                          textColor=colors.HexColor("#475569"))))
    story.append(PageBreak())

    # ---------------------------- table of contents -------------------------
    story.append(Paragraph("Contents", toc_head))
    story.append(Spacer(1, 4))
    toc = TableOfContents()
    toc.levelStyles = [toc1, toc2]
    toc.dotsMinLevel = 0
    story.append(toc)
    story.append(Spacer(1, 14))
    story.append(Paragraph("Appendix A - Project file list", h_small))
    story.append(Paragraph("Appendix B - What changed in this compilation", h_small))
    story.append(PageBreak())

    # ---------------------------- activities --------------------------------
    for entry in ACTIVITIES:
        story.append(Paragraph("Activity %d" % entry["number"], h_act))
        story.append(Paragraph("Activity %d  -  %s" % (entry["number"], entry["title"]),
                               h_act_title))
        story.append(Paragraph("Project : %s&nbsp;&nbsp;&nbsp;|&nbsp;&nbsp;&nbsp;Author : %s"
                               "&nbsp;&nbsp;&nbsp;|&nbsp;&nbsp;&nbsp;%s"
                               % (entry["project"], AUTHOR, COURSE), small))
        story.append(Spacer(1, 6))

        story.append(Paragraph("Description", h_sec))
        for text in entry["description"]:
            story.append(Paragraph(text, body))

        if entry.get("objectives"):
            story.append(Paragraph("Objectives", h_sec))
            for item in entry["objectives"]:
                story.append(Paragraph(item, bullet_style, bulletText="\u2022"))

        for figure_path, caption in figures.get(entry["number"], []):
            from PIL import Image as PILImage
            with PILImage.open(figure_path) as img:
                width_px, height_px = img.size
            target_width = 16.4 * cm
            target_height = target_width * height_px / float(width_px)
            max_height = 11.5 * cm
            if target_height > max_height:
                target_height = max_height
                target_width = target_height * width_px / float(height_px)
            story.append(Spacer(1, 6))
            story.append(Image(figure_path, width=target_width, height=target_height))
            story.append(Paragraph(caption, capt))

        story.append(Paragraph("Source Code", h_sec))
        for label, lines in code_blocks(entry):
            wrapped = []
            for line in lines:
                if len(line) > 108:
                    wrapped.extend(textwrap.wrap(line, 108, subsequent_indent="    ",
                                                 break_long_words=True,
                                                 drop_whitespace=False) or [""])
                else:
                    wrapped.append(line if line else " ")
            story.append(Paragraph(label, code_label))
            story.append(Preformatted("\n".join(wrapped), code_style))

        if entry.get("output"):
            story.append(Paragraph("Sample Output", h_sec))
            story.append(Preformatted("\n".join(line if line else " " for line in entry["output"]),
                                      code_style))

        if entry.get("notes"):
            story.append(Paragraph("Notes", h_sec))
            for note in entry["notes"]:
                story.append(Paragraph(note, body))

        story.append(PageBreak())

    # ---------------------------- appendix A --------------------------------
    story.append(Paragraph("Appendix A - Project file list", h_act))
    story.append(Paragraph("Every folder below is inside the COMPILATION folder of the submitted "
                           "archive. The bin, obj and .suo files of the original projects are not "
                           "included.", small))
    story.append(Spacer(1, 8))
    for project, files in project_files():
        story.append(Paragraph(project, h_small))
        for item in files:
            story.append(Paragraph(item, ParagraphStyle("z", parent=small, fontName="Courier",
                                                        fontSize=7.6, leading=10,
                                                        leftIndent=8)))
    story.append(PageBreak())

    # ---------------------------- appendix B --------------------------------
    story.append(Paragraph("Appendix B - What changed in this compilation", h_act))
    story.append(Spacer(1, 6))
    changes = [
        ("Author and course",
         "The name Alvarez, Rex was replaced by Mangalili, Fhetz and the course line now reads "
         "BSIT 3B in every project, every form title, the About form, the application "
         "configuration file and the console programs."),
        ("Activity 4",
         "New two panel design with a dark banner, a shared ReadInputs() function, four decimal "
         "result formatting and a division by zero guard."),
        ("Activity 5",
         "New side rail layout with three result cards, the Enter key runs the evaluation and "
         "the first card changes colour with the result."),
        ("Activity 6",
         "New left card and right result layout with a large average badge, one ReadGrades() "
         "validation function and a 0 to 100 range check."),
        ("Activity 8",
         "The four overlapping GroupBoxes were replaced by a compose card and three option "
         "cards; the long If chain was replaced by helper functions."),
        ("Activity 10",
         "Both forms were rebuilt: Panel Playground with an event log and an API inspector, and "
         "Split Container Lab with a category list, a record list, a details pane and a reset "
         "layout feature."),
        ("Activity 12",
         "The explorer was rewritten as Mangalili Explorer with history stacks, an address bar, "
         "a live filter, numeric column sorting, safe image preview and layout buttons for the "
         "nested splitters."),
    ]
    for label, text in changes:
        story.append(Paragraph("<b>%s</b>  -  %s" % (label, text),
                               ParagraphStyle("c", parent=body, spaceAfter=6)))
    story.append(Spacer(1, 18))
    story.append(Paragraph("Prepared by : %s        Course : %s        Date : %s"
                           % (AUTHOR, COURSE, DATE_LINE), small))

    doc.multiBuild(story)
    return out_path


def main():
    os.makedirs(DOCS, exist_ok=True)
    figures = build_figures()
    docx_path = os.path.join(DOCS, DOC_BASENAME + ".docx")
    pdf_path = os.path.join(DOCS, DOC_BASENAME + ".pdf")
    build_docx(figures, docx_path)
    print("wrote", docx_path)
    build_pdf(figures, pdf_path)
    print("wrote", pdf_path)


if __name__ == "__main__":
    main()
