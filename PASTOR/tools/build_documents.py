"""Build the compilation documents (DOCX + PDF) of the student.

Contents of the document - only these four parts per activity:
    1. Title of the activity
    2. Form Design  (picture of the designer layout, form activities only)
    3. Source Code  (the real source files)
    4. Sample Output (picture of the running form, or the console text)

No table of contents, no description, no objectives, no notes, no header and no
footer, no appendices.

Run from the repository root:
    python3 tools/build_documents.py

Writes:
    docs/<Name>_Activity_Compilation.docx
    docs/<Name>_Activity_Compilation.pdf
    docs/figures/*.png
"""
import os
import sys
import textwrap

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from activities import (ACTIVITIES, AUTHOR, COURSE, SUBJECT, SCHOOL_YEAR, DATE_LINE)
from design_figure import render as render_figure
from samples import SAMPLES

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
COMPILATION = os.path.join(ROOT, "COMPILATION")
DOCS = os.path.join(ROOT, "docs")
FIGURES = os.path.join(DOCS, "figures")
PHOTO_PATH = os.path.join(COMPILATION, "IntegProgProject2026", "IntegProgProject2026", "Images", "photo.jpg")

SURNAME = AUTHOR.split(",")[0].strip()
GIVEN = AUTHOR.split(",")[1].strip().replace(" ", "")
DOC_BASENAME = "%s_%s_Activity_Compilation" % (SURNAME, GIVEN)

INK = (0x06, 0x2C, 0x2A)
ACCENT = (0x0C, 0x4A, 0x45)
MUTED = (0x5B, 0x7C, 0x78)
CODE_BG = "F2F5F9"


# ---------------------------------------------------------------------------
# source code
# ---------------------------------------------------------------------------
def read_source(relative_path):
    path = os.path.join(COMPILATION, relative_path)
    with open(path, "r", encoding="utf-8-sig", errors="replace") as handle:
        text = handle.read()
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", "    ")
    lines = []
    skipped = False
    for line in text.split("\n"):
        if len(line) > 320:
            if not skipped:
                lines.append(line[:120].rstrip() + " ...")
                lines.append("' [very long data line omitted in this document]")
                skipped = True
            continue
        skipped = False
        lines.append(line.rstrip())
    while lines and lines[-1] == "":
        lines.pop()
    return lines


def code_blocks(entry):
    """Only the plain .vb source files.

    The designer code (`*.Designer.vb`), the `My Project` files, the project
    file and the solution file are NOT printed: the design is already shown as
    a picture (Form Design) and the sample output as another picture.
    """
    return code_blocks_from(entry["files"])


def code_blocks_from(files):
    blocks = []
    for path, label in files:
        low = path.replace("\\", "/").lower()
        if not low.endswith(".vb"):
            continue
        if low.endswith(".designer.vb") or "my project/" in low:
            continue
        blocks.append((label, read_source(path)))
    return blocks


# ---------------------------------------------------------------------------
# pictures
# ---------------------------------------------------------------------------
def fill(value):
    """Replace {USER} / {AUTHOR} / {COURSE} tokens inside the sample values."""
    if isinstance(value, str):
        return value.format(USER=SURNAME, AUTHOR=AUTHOR, COURSE=COURSE, PHOTO=PHOTO_PATH)
    if isinstance(value, dict):
        return dict((key, fill(item)) for key, item in value.items())
    if isinstance(value, (list, tuple)):
        return [fill(item) for item in value]
    return value


TOOLS_DIR = os.path.dirname(os.path.abspath(__file__))


def resolve_designer(path):
    """`tools:` paths are picture-only layouts kept next to the scripts."""
    if path.startswith("tools:"):
        return os.path.join(TOOLS_DIR, path[len("tools:"):])
    return os.path.join(COMPILATION, path)


def build_figures():
    """Render every picture and return {activity number: [section, ...]}.

    A section is one form: its Form Design picture, its source file(s) and its
    Sample Output picture.  Single-form activities have exactly one section
    without a name; activities with several forms (10, 11, 13) have one
    section per form, in the order of `GROUPS` in activities.py.
    """
    os.makedirs(FIGURES, exist_ok=True)
    for name in os.listdir(FIGURES):
        if name.lower().endswith(".png"):
            os.remove(os.path.join(FIGURES, name))

    sections = {}
    for entry in ACTIVITIES:
        number = entry["number"]
        samples = SAMPLES.get(number, [])
        result = []

        groups = entry.get("groups")
        if groups:
            for index, (name, files, designer) in enumerate(groups, 1):
                section = {"name": name, "designs": [], "outputs": [],
                           "code": code_blocks_from(files), "text": None}
                if designer:
                    out_path = os.path.join(FIGURES, "design_act%02d_%d.png" % (number, index))
                    render_figure(resolve_designer(designer), out_path, mode="design")
                    section["designs"].append((out_path, "Form design - %s" % name.split("  -  ")[0]))
                    for sample in samples:
                        if sample["designer"] == designer:
                            out_path = os.path.join(FIGURES, "output_act%02d_%d.png" % (number, index))
                            render_figure(resolve_designer(designer), out_path,
                                          overrides=fill(sample.get("values")))
                            section["outputs"].append((out_path, sample["caption"]))
                            break
                result.append(section)
        else:
            section = {"name": None, "designs": [], "outputs": [],
                       "code": code_blocks(entry), "text": None}
            for designer_path, caption in entry.get("figures", []):
                out_path = os.path.join(FIGURES, "design_act%02d_%d.png"
                                        % (number, len(section["designs"]) + 1))
                render_figure(resolve_designer(designer_path), out_path, mode="design")
                section["designs"].append((out_path, caption))
            for sample in samples:
                out_path = os.path.join(FIGURES, "output_act%02d_%d.png"
                                        % (number, len(section["outputs"]) + 1))
                render_figure(resolve_designer(sample["designer"]), out_path,
                              overrides=fill(sample.get("values")))
                section["outputs"].append((out_path, sample["caption"]))
            if not section["outputs"]:
                section["text"] = entry.get("output", [])
            result.append(section)
        sections[number] = result

    return sections


# ---------------------------------------------------------------------------
# shared helpers
# ---------------------------------------------------------------------------
def wrap_code(lines, width=104):
    out = []
    for line in lines:
        if len(line) > width:
            out.extend(textwrap.wrap(line, width, subsequent_indent="    ",
                                     break_long_words=True, drop_whitespace=False) or [""])
        else:
            out.append(line)
    return out


def activity_heading(entry):
    return "Activity %d - %s" % (entry["number"], entry["title"])


def activity_meta(entry):
    return "%s   |   %s   |   %s" % (entry["project"], AUTHOR, COURSE)


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def build_docx(sections, out_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Cm, RGBColor

    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)

    for name, size, colour in (("Heading 1", 19, INK),
                               ("Heading 2", 12.5, ACCENT),
                               ("Heading 3", 11, MUTED)):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(*colour)
        style.font.bold = True

    def shade(paragraph, fill):
        properties = paragraph._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        properties.append(shd)

    def paragraph(text, size=10.5, bold=False, italic=False, colour=None,
                  align=None, space_after=4, space_before=0, style=None):
        par = doc.add_paragraph(style=style)
        if align == "center":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(space_after)
        par.paragraph_format.space_before = Pt(space_before)
        run = par.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.font.name = "Calibri"
        if colour:
            run.font.color.rgb = RGBColor(*colour)
        return par

    def code_block(lines, label=None):
        if label:
            par = doc.add_paragraph()
            par.paragraph_format.space_before = Pt(6)
            par.paragraph_format.space_after = Pt(2)
            run = par.add_run(label)
            run.bold = True
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(*ACCENT)
        for line in lines:
            par = doc.add_paragraph()
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.space_before = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.left_indent = Cm(0.2)
            shade(par, CODE_BG)
            run = par.add_run(line if line else " ")
            run.font.name = "Consolas"
            run.font.size = Pt(7.5)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def picture(path, width_cm=16.6):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(4)
        par.paragraph_format.space_after = Pt(8)
        par.add_run().add_picture(path, width=Cm(width_cm))

    # ------------------------------ title page ------------------------------
    for _ in range(4):
        doc.add_paragraph()
    paragraph(SUBJECT.upper(), 12, True, colour=MUTED, align="center", space_after=10)
    paragraph("ACTIVITY COMPILATION", 30, True, colour=INK, align="center", space_after=6)
    paragraph("Activities 1 to %d" % len(ACTIVITIES), 13, colour=MUTED,
              align="center", space_after=40)
    paragraph(AUTHOR, 20, True, colour=ACCENT, align="center", space_after=4)
    paragraph(COURSE, 13, colour=MUTED, align="center", space_after=40)
    paragraph("Subject  :  %s" % SUBJECT, 11, align="center", space_after=4)
    paragraph("School Year  :  %s" % SCHOOL_YEAR, 11, align="center", space_after=4)
    paragraph("Date Submitted  :  %s" % DATE_LINE, 11, align="center", space_after=4)
    paragraph("Submitted to  :  ____________________", 11, align="center", space_after=4)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------- activities -----------------------------
    for index, entry in enumerate(ACTIVITIES):
        number = entry["number"]

        paragraph(activity_heading(entry), 19, True, colour=INK, space_after=2)
        paragraph(activity_meta(entry), 9.5, colour=MUTED, space_after=10)

        for part in sections[number]:
            if part["name"]:
                paragraph(part["name"], 14, True, colour=INK, space_before=12, space_after=2)
            if part["designs"]:
                paragraph("Form Design", 12.5, True, colour=ACCENT, space_before=4, space_after=4)
                for path, caption in part["designs"]:
                    picture(path)
                    if caption:
                        paragraph(caption, 9, italic=True, colour=MUTED,
                                  align="center", space_after=10)

            paragraph("Source Code", 12.5, True, colour=ACCENT, space_before=6, space_after=4)
            for label, lines in part["code"]:
                code_block(wrap_code(lines), label)

            if part["outputs"] or part["text"] is not None:
                paragraph("Sample Output", 12.5, True, colour=ACCENT, space_before=8, space_after=4)
                for path, caption in part["outputs"]:
                    picture(path)
                    if caption:
                        paragraph(caption, 9, italic=True, colour=MUTED,
                                  align="center", space_after=10)
                if part["text"] is not None:
                    code_block(wrap_code(part["text"]))

        if index != len(ACTIVITIES) - 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    doc.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def build_pdf(sections, out_path):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (Image, KeepTogether, PageBreak, Paragraph,
                                    Preformatted, SimpleDocTemplate, Spacer)
    from PIL import Image as PILImage

    def register(name, path):
        try:
            pdfmetrics.registerFont(TTFont(name, path))
            return name
        except Exception:
            return "Helvetica"

    body_font = register("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    try:
        pdfmetrics.registerFontFamily(body_font, normal=body_font, bold=body_font,
                                      italic=body_font, boldItalic=body_font)
    except Exception:
        pass
    mono_font = register("DejaVuMono", "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf")

    title_style = ParagraphStyle("title", fontName=body_font, fontSize=27, leading=33,
                                 textColor=colors.Color(INK[0] / 255.0, INK[1] / 255.0, INK[2] / 255.0),
                                 alignment=TA_CENTER, spaceAfter=12)
    big_style = ParagraphStyle("big", parent=title_style, fontSize=17, leading=22)
    meta_style = ParagraphStyle("meta", fontName=body_font, fontSize=9.5, leading=14,
                                textColor=colors.Color(MUTED[0] / 255.0, MUTED[1] / 255.0, MUTED[2] / 255.0),
                                spaceAfter=8)
    h1 = ParagraphStyle("h1", fontName=body_font, fontSize=17, leading=21,
                        textColor=colors.Color(INK[0] / 255.0, INK[1] / 255.0, INK[2] / 255.0),
                        spaceBefore=2, spaceAfter=4)
    h2 = ParagraphStyle("h2", fontName=body_font, fontSize=11.5, leading=15,
                        textColor=colors.Color(ACCENT[0] / 255.0, ACCENT[1] / 255.0, ACCENT[2] / 255.0),
                        spaceBefore=12, spaceAfter=5)
    caption_style = ParagraphStyle("cap", fontName=body_font, fontSize=8.5, leading=11,
                                   textColor=colors.Color(MUTED[0] / 255.0, MUTED[1] / 255.0, MUTED[2] / 255.0),
                                   alignment=TA_CENTER, spaceAfter=10)
    code_style = ParagraphStyle("code", fontName=mono_font, fontSize=6.4, leading=7.7,
                                backColor=colors.HexColor("#" + CODE_BG),
                                borderPadding=(2, 3, 2, 3), leftIndent=6, spaceAfter=0)

    max_width = 16.6 * cm

    def picture(path):
        with PILImage.open(path) as raw:
            width, height = raw.size
        scale = min(max_width / width, (20.5 * cm) / height)
        return Image(path, width=width * scale, height=height * scale, hAlign="CENTER")

    story = []
    story.append(Spacer(1, 6.5 * cm))
    story.append(Paragraph(SUBJECT.upper(), meta_style))
    story.append(Paragraph("ACTIVITY COMPILATION", title_style))
    story.append(Spacer(1, 6))
    story.append(Paragraph("Activities 1 to %d" % len(ACTIVITIES), big_style))
    story.append(Spacer(1, 2.6 * cm))
    story.append(Paragraph(AUTHOR, ParagraphStyle("author", parent=big_style, fontSize=19,
                                                  textColor=colors.Color(ACCENT[0] / 255.0,
                                                                         ACCENT[1] / 255.0,
                                                                         ACCENT[2] / 255.0))))
    story.append(Paragraph(COURSE, ParagraphStyle("course", parent=meta_style, fontSize=12)))
    story.append(Spacer(1, 2.4 * cm))
    for line in ("Subject  :  %s" % SUBJECT,
                 "School Year  :  %s" % SCHOOL_YEAR,
                 "Date Submitted  :  %s" % DATE_LINE,
                 "Submitted to  :  ____________________"):
        story.append(Paragraph(line, meta_style))
    story.append(PageBreak())

    for index, entry in enumerate(ACTIVITIES):
        number = entry["number"]
        story.append(Paragraph(activity_heading(entry), h1))
        story.append(Paragraph(activity_meta(entry), meta_style))

        for part in sections[number]:
            if part["name"]:
                story.append(Paragraph(part["name"], ParagraphStyle(
                    "grp", fontName=body_font, fontSize=13.5, leading=17, spaceBefore=14,
                    spaceAfter=2, textColor=colors.Color(INK[0] / 255.0, INK[1] / 255.0,
                                                         INK[2] / 255.0))))
            if part["designs"]:
                first = True
                for path, caption in part["designs"]:
                    block = []
                    if first:
                        block.append(Paragraph("Form Design", h2))
                        first = False
                    block.append(picture(path))
                    if caption:
                        block.append(Paragraph(caption, caption_style))
                    story.append(KeepTogether(block))

            story.append(Paragraph("Source Code", h2))
            for label, lines in part["code"]:
                story.append(Paragraph(label, ParagraphStyle("lbl", parent=meta_style,
                                                             fontSize=8.5, spaceBefore=4,
                                                             spaceAfter=2)))
                story.append(Preformatted("\n".join(wrap_code(lines, 112)), code_style))
                story.append(Spacer(1, 6))

            if part["outputs"]:
                first = True
                for path, caption in part["outputs"]:
                    block = []
                    if first:
                        block.append(Paragraph("Sample Output", h2))
                        first = False
                    block.append(picture(path))
                    if caption:
                        block.append(Paragraph(caption, caption_style))
                    story.append(KeepTogether(block))
            elif part["text"] is not None:
                story.append(Paragraph("Sample Output", h2))
                story.append(Preformatted("\n".join(wrap_code(part["text"], 112)), code_style))

        if index != len(ACTIVITIES) - 1:
            story.append(PageBreak())

    doc = SimpleDocTemplate(out_path, pagesize=A4,
                            leftMargin=2.0 * cm, rightMargin=2.0 * cm,
                            topMargin=1.8 * cm, bottomMargin=1.8 * cm,
                            title="%s - Activity Compilation" % SURNAME,
                            author=AUTHOR, subject=SUBJECT)
    doc.build(story)
    return out_path


# ---------------------------------------------------------------------------
def main():
    os.makedirs(DOCS, exist_ok=True)
    sections = build_figures()
    docx_path = os.path.join(DOCS, DOC_BASENAME + ".docx")
    pdf_path = os.path.join(DOCS, DOC_BASENAME + ".pdf")
    build_docx(sections, docx_path)
    print("wrote", docx_path)
    build_pdf(sections, pdf_path)
    print("wrote", pdf_path)


if __name__ == "__main__":
    main()
